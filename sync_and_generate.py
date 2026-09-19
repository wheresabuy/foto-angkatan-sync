#!/usr/bin/env python3
"""
Sync and Document Generator for Pas Foto 3x4 Teknik Komputer 2025
================================================================
Memantau Google Drive angkatan secara otomatis, mengunduh foto baru,
memotong ke format 3x4 cm (300 DPI), meregenerasi dokumen cetak PDF & Word,
serta mengunggah hasil terbaru kembali ke Google Drive.

Dapat berjalan di lokal (Ubuntu/GNOME) maupun di cloud (GitHub Actions).
"""

import argparse
import json
import os
import re
import subprocess
import sys
import time
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from PIL import Image, ImageOps, ImageDraw, ImageFont

# Optional docx and reportlab imports
try:
    import docx
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except ImportError:
    print("Warning: python-docx not installed. Run: pip install python-docx")

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.pdfgen import canvas
    from reportlab.lib.units import cm
except ImportError:
    print("Warning: reportlab not installed. Run: pip install reportlab")

# Paths Configuration
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
RAW_IMG_DIR = os.path.join(DATA_DIR, "raw_images")
PROC_IMG_DIR = os.path.join(DATA_DIR, "processed_3x4")
STATE_FILE = os.path.join(DATA_DIR, "students_state.json")
BIO_FILE = os.path.join(DATA_DIR, "biodata.json")
ENV_FILE = os.path.join(BASE_DIR, ".env.github")

# Google Drive Target Folder ID
DRIVE_ROOT_FOLDER_ID = "1HS4EInIOlsWEXKPP05sSapKE1Dt0G72W"

# Image & Page Layout Constants
TARGET_W, TARGET_H = 354, 472  # 3x4 cm @ 300 DPI
COLS = 5
ROWS = 5
ITEMS_PER_PAGE = COLS * ROWS

def get_access_token():
    """
    Mengambil OAuth2 token:
    1. Dari environment variable (GDRIVE_REFRESH_TOKEN, dll.) -> Cocok untuk GitHub Actions.
    2. Dari file lokal .env.github (jika ada).
    3. Dari GNOME Online Accounts via D-Bus session -> Cocok untuk laptop lokal.
    """
    # 1. Cek Environment Variables atau .env.github
    refresh_token = os.environ.get("GDRIVE_REFRESH_TOKEN")
    client_id = os.environ.get("GDRIVE_CLIENT_ID")
    client_secret = os.environ.get("GDRIVE_CLIENT_SECRET")

    if not refresh_token and os.path.exists(ENV_FILE):
        try:
            with open(ENV_FILE) as f:
                for line in f:
                    if "=" in line:
                        k, v = line.strip().split("=", 1)
                        if k == "GDRIVE_REFRESH_TOKEN" and not refresh_token:
                            refresh_token = v
                        elif k == "GDRIVE_CLIENT_ID" and not client_id:
                            client_id = v
                        elif k == "GDRIVE_CLIENT_SECRET" and not client_secret:
                            client_secret = v
        except Exception:
            pass

    if refresh_token and client_id and client_secret:
        try:
            data = urllib.parse.urlencode({
                "client_id": client_id,
                "client_secret": client_secret,
                "refresh_token": refresh_token,
                "grant_type": "refresh_token"
            }).encode("utf-8")
            req = urllib.request.Request(
                "https://oauth2.googleapis.com/token",
                data=data,
                headers={"Content-Type": "application/x-www-form-urlencoded"}
            )
            with urllib.request.urlopen(req) as resp:
                token_data = json.loads(resp.read().decode("utf-8"))
                return token_data.get("access_token")
        except Exception as e:
            print(f"[WARN] Refresh token OAuth2 gagal: {e}")

    # 2. Fallback ke GNOME Online Accounts (Local DBus)
    cmd = [
        "gdbus", "call", "--session", "--dest", "org.gnome.OnlineAccounts",
        "--object-path", "/org/gnome/OnlineAccounts/Accounts/account_1789698921_0",
        "--method", "org.gnome.OnlineAccounts.OAuth2Based.GetAccessToken"
    ]
    try:
        output = subprocess.check_output(cmd, stderr=subprocess.DEVNULL).decode()
        return output.split("'")[1]
    except Exception:
        pass

    print("[ERROR] Tidak dapat memperoleh access token Google Drive.")
    return None

def fetch_drive_catalog(token):
    """Mengambil seluruh daftar subfolder dan file di dalamnya dari Google Drive."""
    all_folders = []
    page_token = None
    while True:
        q = f'"{DRIVE_ROOT_FOLDER_ID}" in parents and trashed = false'
        url = f"https://www.googleapis.com/drive/v3/files?q={urllib.parse.quote(q)}&fields=nextPageToken,files(id,name,mimeType,size,modifiedTime)&pageSize=100"
        if page_token:
            url += f"&pageToken={page_token}"
        req = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
        with urllib.request.urlopen(req) as resp:
            res = json.loads(resp.read().decode("utf-8"))
            all_folders.extend(res.get("files", []))
            page_token = res.get("nextPageToken")
            if not page_token:
                break

    def list_subfolder(folder):
        f_id = folder["id"]
        files = []
        pt = None
        while True:
            q = f'"{f_id}" in parents and trashed = false'
            url = f"https://www.googleapis.com/drive/v3/files?q={urllib.parse.quote(q)}&fields=nextPageToken,files(id,name,mimeType,size,modifiedTime)&pageSize=100"
            if pt:
                url += f"&pageToken={pt}"
            req = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
            try:
                with urllib.request.urlopen(req) as resp:
                    res = json.loads(resp.read().decode("utf-8"))
                    files.extend(res.get("files", []))
                    pt = res.get("nextPageToken")
                    if not pt:
                        break
            except Exception as e:
                print(f"[WARN] Gagal membaca subfolder {folder['name']}: {e}")
                break
        return {"folder": folder, "files": files}

    with ThreadPoolExecutor(max_workers=10) as ex:
        catalog = list(ex.map(list_subfolder, all_folders))

    return catalog

def select_best_photo(short_nrp, img_files):
    """Memilih foto terbaik jika ada lebih dari 1 file gambar."""
    if not img_files:
        return None
    if len(img_files) == 1:
        return img_files[0]
    if "027" in short_nrp:
        for f in img_files:
            if "IMG20260907145531" in f["name"]:
                return f
    elif "110" in short_nrp:
        for f in img_files:
            if "IMG-20260919-WA0026" in f["name"]:
                return f
    return img_files[0]

def download_file(token, file_info, target_path):
    """Mengunduh gambar resolusi tinggi dari Google Drive."""
    fid = file_info["id"]
    try:
        url = f"https://www.googleapis.com/drive/v3/files/{fid}?fields=id,name,thumbnailLink"
        req = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
        with urllib.request.urlopen(req) as resp:
            meta = json.loads(resp.read().decode())
        thumb = meta.get("thumbnailLink")
        if thumb:
            dl_url = thumb.replace("=s220", "=s1600")
            req_dl = urllib.request.Request(dl_url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req_dl) as resp_dl, open(target_path, "wb") as f_out:
                f_out.write(resp_dl.read())
            return True
    except Exception:
        pass

    try:
        url = f"https://www.googleapis.com/drive/v3/files/{fid}?alt=media"
        req = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
        with urllib.request.urlopen(req) as resp, open(target_path, "wb") as f_out:
            f_out.write(resp.read())
        return True
    except Exception as e:
        print(f"[ERROR] Download direct {file_info['name']} gagal: {e}")
        return False

def crop_and_frame_3x4(raw_path, out_path, short_nrp):
    """Memotong foto ke rasio 3:4 dengan framing wajah optimal dan border abu-abu."""
    try:
        im = Image.open(raw_path)
        im = ImageOps.exif_transpose(im).convert("RGB")
        
        center_x = 0.75 if short_nrp == "060" else 0.5
        im_crop = ImageOps.fit(im, (TARGET_W, TARGET_H), centering=(center_x, 0.42))
        
        draw = ImageDraw.Draw(im_crop)
        draw.rectangle([(0, 0), (TARGET_W - 1, TARGET_H - 1)], outline=(200, 200, 200), width=1)
        
        im_crop.save(out_path, "JPEG", quality=95, dpi=(300, 300))
        return True
    except Exception as e:
        print(f"[ERROR] Crop {short_nrp} gagal: {e}")
        return False

def create_placeholder(out_path, short_nrp):
    """Membuat placeholder garis putus-putus untuk teman yang belum upload foto."""
    im = Image.new("RGB", (TARGET_W, TARGET_H), color=(245, 246, 248))
    draw = ImageDraw.Draw(im)
    border_color = (170, 175, 185)
    dash_len = 8
    for x in range(0, TARGET_W, dash_len * 2):
        draw.line([(x, 0), (min(x + dash_len, TARGET_W - 1), 0)], fill=border_color, width=2)
        draw.line([(x, TARGET_H - 1), (min(x + dash_len, TARGET_W - 1), TARGET_H - 1)], fill=border_color, width=2)
    for y in range(0, TARGET_H, dash_len * 2):
        draw.line([(0, y), (0, min(y + dash_len, TARGET_H - 1))], fill=border_color, width=2)
        draw.line([(TARGET_W - 1, y), (TARGET_W - 1, min(y + dash_len, TARGET_H - 1))], fill=border_color, width=2)
    
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 20)
        font_sub = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 16)
    except Exception:
        font = ImageFont.load_default()
        font_sub = font

    draw.text((TARGET_W//2, TARGET_H//2 - 30), "FOTO 3x4", fill=(130, 135, 145), font=font, anchor="mm")
    draw.text((TARGET_W//2, TARGET_H//2 + 5), f"NRP: {short_nrp}", fill=(100, 105, 115), font=font_sub, anchor="mm")
    draw.text((TARGET_W//2, TARGET_H//2 + 35), "[ Belum Upload ]", fill=(160, 165, 175), font=font_sub, anchor="mm")
    
    im.save(out_path, "JPEG", quality=95, dpi=(300, 300))

# ==================== DOKUMEN GENERATOR ====================

def generate_pdf_documents(students, output_path, title, subtitle):
    """Menghasilkan dokumen PDF A4 siap cetak dengan ReportLab."""
    c = canvas.Canvas(output_path, pagesize=A4)
    PAGE_W, PAGE_H = A4
    PHOTO_W = 3.0 * cm
    PHOTO_H = 4.0 * cm
    LABEL_GAP = 0.12 * cm
    LINE_HEIGHT = 0.32 * cm
    COL_GAP = 0.55 * cm
    ROW_GAP = 0.35 * cm
    GRID_W = COLS * PHOTO_W + (COLS - 1) * COL_GAP
    LEFT_MARGIN = (PAGE_W - GRID_W) / 2.0

    def truncate_text(text, font_name, font_size, max_w):
        c.setFont(font_name, font_size)
        if c.stringWidth(text, font_name, font_size) <= max_w:
            return text
        while len(text) > 3 and c.stringWidth(text + "...", font_name, font_size) > max_w:
            text = text[:-1]
        return text + "..."

    total_pages = (len(students) + ITEMS_PER_PAGE - 1) // ITEMS_PER_PAGE

    for page_idx in range(total_pages):
        page_num = page_idx + 1
        c.saveState()
        c.setFont("Helvetica-Bold", 10.5)
        c.setFillColor(colors.HexColor("#1A365D"))
        c.drawCentredString(PAGE_W / 2.0, PAGE_H - 1.1 * cm, title)
        c.setFont("Helvetica", 7.5)
        c.setFillColor(colors.HexColor("#4A5568"))
        c.drawCentredString(PAGE_W / 2.0, PAGE_H - 1.5 * cm, f"{subtitle}  |  Halaman {page_num} dari {total_pages}")
        c.setStrokeColor(colors.HexColor("#CBD5E0"))
        c.setLineWidth(0.6)
        c.line(LEFT_MARGIN, PAGE_H - 1.65 * cm, PAGE_W - LEFT_MARGIN, PAGE_H - 1.65 * cm)
        c.restoreState()

        page_students = students[page_idx * ITEMS_PER_PAGE : (page_idx + 1) * ITEMS_PER_PAGE]
        for item_idx, s in enumerate(page_students):
            col = item_idx % COLS
            row = item_idx // COLS
            x = LEFT_MARGIN + col * (PHOTO_W + COL_GAP)
            top_y = PAGE_H - 1.85 * cm
            cell_total_h = PHOTO_H + LABEL_GAP + 2 * LINE_HEIGHT
            y_photo = top_y - row * (cell_total_h + ROW_GAP) - PHOTO_H

            nrp = s["short_nrp"]
            img_p = os.path.join(PROC_IMG_DIR, f"{nrp}_3x4.jpg")
            if os.path.exists(img_p):
                c.drawImage(img_p, x, y_photo, width=PHOTO_W, height=PHOTO_H, preserveAspectRatio=True)

            c.setStrokeColor(colors.HexColor("#A0AEC0") if not s["has_photo"] else colors.HexColor("#CBD5E0"))
            c.setLineWidth(0.5)
            c.rect(x, y_photo, PHOTO_W, PHOTO_H, stroke=1, fill=0)

            y_nrp = y_photo - LABEL_GAP - LINE_HEIGHT + 2
            y_name = y_nrp - LINE_HEIGHT

            c.setFont("Helvetica-Bold", 7.5)
            c.setFillColor(colors.HexColor("#2D3748"))
            c.drawCentredString(x + PHOTO_W / 2.0, y_nrp, s.get("full_nrp") or f"5024251{int(nrp):03d}")

            c.setFont("Helvetica", 6.5)
            c.setFillColor(colors.HexColor("#4A5568") if s["has_photo"] else colors.HexColor("#A0AEC0"))
            c.drawCentredString(x + PHOTO_W / 2.0, y_name, truncate_text(s["nama"], "Helvetica", 6.5, PHOTO_W + 0.3 * cm))

        c.showPage()

    c.save()

def generate_docx_documents(students, output_path, title_text, subtitle_text):
    """Menghasilkan dokumen Word .docx A4 dengan python-docx."""
    doc = docx.Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)

    total_pages = (len(students) + ITEMS_PER_PAGE - 1) // ITEMS_PER_PAGE

    def set_cell_margins(cell, top=40, bottom=40, left=40, right=40):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    for page_idx in range(total_pages):
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(0)
        p_title.paragraph_format.space_after = Pt(2)
        r_title = p_title.add_run(title_text)
        r_title.font.name = 'Arial'
        r_title.font.size = Pt(10.5)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(26, 54, 93)

        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(6)
        r_sub = p_sub.add_run(f"{subtitle_text}  |  Halaman {page_idx + 1} dari {total_pages}")
        r_sub.font.name = 'Arial'
        r_sub.font.size = Pt(7.5)
        r_sub.font.color.rgb = RGBColor(74, 85, 104)

        page_items = students[page_idx * ITEMS_PER_PAGE : (page_idx + 1) * ITEMS_PER_PAGE]
        num_rows = (len(page_items) + COLS - 1) // COLS
        table = doc.add_table(rows=num_rows, cols=COLS)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        col_width = Cm(3.6)
        for r_idx in range(num_rows):
            row = table.rows[r_idx]
            for c_idx in range(COLS):
                cell = row.cells[c_idx]
                cell.width = col_width
                cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                set_cell_margins(cell)

                item_idx = r_idx * COLS + c_idx
                if item_idx < len(page_items):
                    s = page_items[item_idx]
                    nrp = s["short_nrp"]
                    img_path = os.path.join(PROC_IMG_DIR, f"{nrp}_3x4.jpg")

                    p_img = cell.paragraphs[0]
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(0)
                    p_img.paragraph_format.space_after = Pt(2)

                    if os.path.exists(img_path):
                        p_img.add_run().add_picture(img_path, width=Cm(3.0), height=Cm(4.0))

                    p_nrp = cell.add_paragraph()
                    p_nrp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_nrp.paragraph_format.space_before = Pt(1)
                    p_nrp.paragraph_format.space_after = Pt(0)
                    r_nrp = p_nrp.add_run(s.get("full_nrp") or f"5024251{int(nrp):03d}")
                    r_nrp.font.name = 'Arial'
                    r_nrp.font.size = Pt(7.5)
                    r_nrp.font.bold = True
                    r_nrp.font.color.rgb = RGBColor(45, 55, 72)

                    p_name = cell.add_paragraph()
                    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_name.paragraph_format.space_before = Pt(0)
                    p_name.paragraph_format.space_after = Pt(4)
                    r_name = p_name.add_run(s["nama"])
                    r_name.font.name = 'Arial'
                    r_name.font.size = Pt(6.5)
                    r_name.font.color.rgb = RGBColor(74, 85, 104) if s["has_photo"] else RGBColor(160, 174, 192)
                else:
                    cell.paragraphs[0].text = ""

        if page_idx < total_pages - 1:
            doc.add_page_break()

    doc.save(output_path)

def upload_file_to_drive(token, local_path, file_name, mime_type):
    """Mengunggah atau mengupdate file langsung di folder root Google Drive."""
    try:
        q = f'"{DRIVE_ROOT_FOLDER_ID}" in parents and name = "{file_name}" and trashed = false'
        url = f"https://www.googleapis.com/drive/v3/files?q={urllib.parse.quote(q)}&fields=files(id)"
        req = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
        with urllib.request.urlopen(req) as resp:
            files = json.loads(resp.read().decode("utf-8")).get("files", [])

        with open(local_path, "rb") as f:
            content = f.read()

        if files:
            fid = files[0]["id"]
            upload_url = f"https://www.googleapis.com/upload/drive/v3/files/{fid}?uploadType=media"
            req_up = urllib.request.Request(
                upload_url, data=content,
                headers={"Authorization": f"Bearer {token}", "Content-Type": mime_type},
                method="PATCH"
            )
            with urllib.request.urlopen(req_up):
                pass
            print(f"[DRIVE] Berhasil mengupdate '{file_name}' di Google Drive.")
        else:
            boundary = "-------314159265358979323846"
            meta = json.dumps({"name": file_name, "parents": [DRIVE_ROOT_FOLDER_ID]})
            body = (
                f"--{boundary}\r\n"
                f"Content-Type: application/json; charset=UTF-8\r\n\r\n"
                f"{meta}\r\n"
                f"--{boundary}\r\n"
                f"Content-Type: {mime_type}\r\n\r\n"
            ).encode() + content + f"\r\n--{boundary}--\r\n".encode()

            upload_url = "https://www.googleapis.com/upload/drive/v3/files?uploadType=multipart"
            req_up = urllib.request.Request(
                upload_url, data=body,
                headers={"Authorization": f"Bearer {token}", "Content-Type": f"multipart/related; boundary={boundary}"},
                method="POST"
            )
            with urllib.request.urlopen(req_up):
                pass
            print(f"[DRIVE] Berhasil mengunggah '{file_name}' baru ke Google Drive.")
        return True
    except Exception as e:
        print(f"[ERROR] Gagal mengunggah {file_name} ke Google Drive: {e}")
        return False

def send_notification(title, message):
    """Mengirim notifikasi desktop Ubuntu/GNOME menggunakan notify-send."""
    try:
        subprocess.run(["notify-send", "-a", "Foto 3x4 Angkatan", title, message], check=False)
    except Exception:
        pass

# ==================== MAIN SYNC PIPELINE ====================

def run_sync_cycle(force_regen=False, upload_drive=False):
    """Menjalankan satu siklus sinkronisasi dan rebuild dokumen jika ada perubahan."""
    token = get_access_token()
    if not token:
        print("[ERROR] Token tidak tersedia.")
        return False

    print(f"[{time.strftime('%H:%M:%S')}] Memeriksa Google Drive...")
    catalog = fetch_drive_catalog(token)
    print(f"[{time.strftime('%H:%M:%S')}] Ditemukan {len(catalog)} folder mahasiswa.")

    bios = {}
    if os.path.exists(BIO_FILE):
        with open(BIO_FILE) as f:
            bios = json.load(f)

    old_state = {}
    if os.path.exists(STATE_FILE):
        with open(STATE_FILE) as f:
            for s in json.load(f):
                old_state[s["short_nrp"]] = s

    updated_students = []
    new_photos_count = 0
    newly_added_names = []

    for item in catalog:
        fol = item["folder"]["name"]
        parts = fol.split(" - ", 1)
        short_nrp = parts[0].strip()
        name_from_fol = parts[1].strip() if len(parts) > 1 else fol

        bio = bios.get(fol, {})
        full_nrp = bio.get("nrp") or (f"5024251{int(short_nrp):03d}" if short_nrp.isdigit() else short_nrp)
        nama = (bio.get("nama") or name_from_fol).strip().title()

        files = item["files"]
        imgs = [fi for fi in files if fi.get("mimeType", "").startswith("image/") or fi["name"].lower().endswith((".jpg", ".jpeg", ".png", ".webp"))]

        best_photo = select_best_photo(short_nrp, imgs)
        has_photo = (best_photo is not None)

        old_s = old_state.get(short_nrp, {})
        old_has_photo = old_s.get("has_photo", False)
        old_photo_file = old_s.get("photo_file")

        if has_photo:
            raw_file_name = f"{short_nrp}_{best_photo['name']}"
            raw_path = os.path.join(RAW_IMG_DIR, raw_file_name)
            proc_path = os.path.join(PROC_IMG_DIR, f"{short_nrp}_3x4.jpg")

            is_new = (not old_has_photo) or (old_photo_file != best_photo["name"]) or (not os.path.exists(proc_path))
            if is_new:
                print(f"[BARU] Mengunduh foto untuk {short_nrp} - {nama} ({best_photo['name']})...")
                download_file(token, best_photo, raw_path)
                crop_and_frame_3x4(raw_path, proc_path, short_nrp)
                new_photos_count += 1
                newly_added_names.append(f"{short_nrp} ({nama.split()[0]})")
        else:
            proc_path = os.path.join(PROC_IMG_DIR, f"{short_nrp}_3x4.jpg")
            if not os.path.exists(proc_path):
                create_placeholder(proc_path, short_nrp)

        updated_students.append({
            "short_nrp": short_nrp,
            "full_nrp": full_nrp,
            "nama": nama,
            "folder": fol,
            "photo_file": best_photo["name"] if best_photo else None,
            "has_photo": has_photo
        })

    updated_students.sort(key=lambda s: int(s["short_nrp"]) if s["short_nrp"].isdigit() else 9999)

    with open(STATE_FILE, "w") as f:
        json.dump(updated_students, f, indent=2)

    total_has = sum(1 for s in updated_students if s["has_photo"])
    print(f"[{time.strftime('%H:%M:%S')}] Total foto tersedia: {total_has}/106.")

    if new_photos_count > 0 or force_regen:
        print(f"[UPDATE] Memperbarui dokumen PDF dan DOCX ({total_has} foto)...")
        photos_only = [s for s in updated_students if s["has_photo"]]

        pdf_gunting = os.path.join(BASE_DIR, "Foto_3x4_Teman_Angkatan_Siap_Gunting.pdf")
        docx_gunting = os.path.join(BASE_DIR, "Foto_3x4_Teman_Angkatan_Siap_Gunting.docx")
        generate_pdf_documents(
            photos_only, pdf_gunting,
            "PAS FOTO 3x4 TEMAN ANGKATAN TEKNIK KOMPUTER 2025",
            f"Siap Cetak & Siap Gunting  •  Urut Sesuai NRP ({total_has} Foto Tersedia)"
        )
        generate_docx_documents(
            photos_only, docx_gunting,
            "PAS FOTO 3x4 TEMAN ANGKATAN TEKNIK KOMPUTER 2025",
            f"Siap Cetak & Siap Gunting  •  Urut Sesuai NRP ({total_has} Foto Tersedia)"
        )

        pdf_lengkap = os.path.join(BASE_DIR, "Foto_3x4_Teman_Angkatan_Lengkap.pdf")
        docx_lengkap = os.path.join(BASE_DIR, "Foto_3x4_Teman_Angkatan_Lengkap.docx")
        generate_pdf_documents(
            updated_students, pdf_lengkap,
            "BUKU ANGKATAN TEKNIK KOMPUTER 2025 - PAS FOTO 3x4",
            "Urut Sesuai NRP  •  Kertas HVS A4 Siap Cetak (Lengkap 106 Mahasiswa)"
        )
        generate_docx_documents(
            updated_students, docx_lengkap,
            "BUKU ANGKATAN TEKNIK KOMPUTER 2025 - PAS FOTO 3x4",
            "Urut Sesuai NRP  •  Kertas HVS A4 Siap Cetak (Lengkap 106 Mahasiswa)"
        )

        print("[SELESAI] Dokumen PDF & Word lokal berhasil diperbarui!")

        if upload_drive:
            print("[DRIVE] Mengunggah dokumen PDF terbaru ke Google Drive angkatan...")
            upload_file_to_drive(token, pdf_gunting, "Foto_3x4_Teman_Angkatan_Siap_Gunting.pdf", "application/pdf")
            upload_file_to_drive(token, pdf_lengkap, "Foto_3x4_Teman_Angkatan_Lengkap.pdf", "application/pdf")

        if new_photos_count > 0:
            names_summary = ", ".join(newly_added_names[:3])
            if len(newly_added_names) > 3:
                names_summary += f" +{len(newly_added_names)-3} lainnya"
            send_notification(
                f"Pas Foto 3x4: {new_photos_count} Foto Baru Masuk!",
                f"Total {total_has}/106 foto tersedia ({names_summary}). Dokumen cetak telah diperbarui otomatis."
            )
        return True
    else:
        print(f"[{time.strftime('%H:%M:%S')}] Tidak ada foto baru. Dokumen sudah versi terbaru.")
        return False

def main():
    parser = argparse.ArgumentParser(description="Auto Sync & Generator Pas Foto 3x4")
    parser.add_argument("--watch", action="store_true", help="Jalankan dalam mode pemantau background")
    parser.add_argument("--interval", type=int, default=300, help="Interval pemantauan dalam detik (default: 300 detik / 5 menit)")
    parser.add_argument("--force-regen", action="store_true", help="Paksa regenerasi dokumen meskipun tidak ada foto baru")
    parser.add_argument("--upload-drive", action="store_true", help="Unggah dokumen PDF hasil regenerasi langsung ke Google Drive")
    args = parser.parse_args()

    os.makedirs(RAW_IMG_DIR, exist_ok=True)
    os.makedirs(PROC_IMG_DIR, exist_ok=True)

    if args.watch:
        print(f"[DAEMON] Mode pemantau aktif. Pengecekan setiap {args.interval} detik...")
        try:
            while True:
                run_sync_cycle(force_regen=args.force_regen, upload_drive=args.upload_drive)
                time.sleep(args.interval)
        except KeyboardInterrupt:
            print("\n[DAEMON] Pemantau dihentikan.")
    else:
        run_sync_cycle(force_regen=args.force_regen, upload_drive=args.upload_drive)

if __name__ == "__main__":
    main()
